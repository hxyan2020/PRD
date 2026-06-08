#!/usr/bin/env python3
"""Generate Murex PFE Design Document (Word) with screenshot placeholders and MLC appendix."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_document_styles(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def add_title_page(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Murex MX.3\nPotential Future Exposure (PFE)\nSolution Design Document")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in [
        "Document Version: 1.0 (Draft)",
        "Status: For Review",
        "Module: Murex Credit Risk / MLC",
        "Classification: Internal",
    ]:
        p = meta.add_run(line + "\n")
        p.font.size = Pt(11)

    doc.add_page_break()


def add_toc_placeholder(doc):
    doc.add_heading("Table of Contents", level=1)
    doc.add_paragraph(
        "Update this section in Word: References → Table of Contents → Automatic Table 1."
    )
    doc.add_page_break()


def add_screenshot_placeholder(doc, caption, notes=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"[INSERT SCREENSHOT: {caption}]")
    run.italic = True
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    # Visual box placeholder
    box = doc.add_paragraph()
    box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    box_run = box.add_run("\n\n\n")
    box_run.font.size = Pt(6)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run(f"Figure: {caption}")
    cap_run.italic = True
    cap_run.font.size = Pt(10)

    if notes:
        note = doc.add_paragraph(notes)
        note.style = "List Bullet"

    doc.add_paragraph()


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_body(doc, text):
    doc.add_paragraph(text)


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_main_document(doc):
    add_title_page(doc)
    add_toc_placeholder(doc)

    # 1. Executive Summary
    add_heading(doc, "1. Executive Summary", 1)
    add_body(
        doc,
        "This document describes the solution design for Potential Future Exposure (PFE) "
        "calculation, monitoring, and limit management within Murex MX.3. PFE is the primary "
        "counterparty credit risk measure used to estimate worst-case exposure over a defined "
        "horizon at a specified confidence level (typically 95% or 97.5%)."
    )
    add_body(
        doc,
        "The design covers methodology selection, data dependencies, batch and real-time "
        "processing, integration with MLC (Murex Limits Controller), reporting, and operational "
        "controls."
    )

    # 2. Scope
    add_heading(doc, "2. Scope", 1)
    add_heading(doc, "2.1 In Scope", 2)
    add_bullets(
        doc,
        [
            "PFE methodology (deterministic add-on, simulation-based, or Monte Carlo)",
            "Counterparty and netting set aggregation",
            "Collateral and CSA impact on exposure profiles",
            "EOD and intraday PFE computation",
            "MLC limit structures, formulae, and breach handling for PFE",
            "Datamart / LRB reporting for credit exposure",
            "Integration with trade lifecycle and pre-deal checks",
        ],
    )
    add_heading(doc, "2.2 Out of Scope", 2)
    add_bullets(
        doc,
        [
            "CVA/DVA/FVA full valuation engine design (referenced only where PFE feeds xVA)",
            "SA-CCR regulatory capital calculation detail (separate design)",
            "Issuer / bond credit risk (unless explicitly in scope for your bank)",
        ],
    )

    # 3. Business Context
    add_heading(doc, "3. Business Context", 1)
    add_body(
        doc,
        "PFE supports internal counterparty limit management, collateral optimisation, "
        "regulatory capital (EAD under IMM or SA-CCR), and stress testing. Limits are "
        "typically set per legal entity, counterparty, product, tenor bucket, or netting set."
    )
    add_screenshot_placeholder(
        doc,
        "PFE business process overview (from design doc)",
        "Replace with your design-doc screenshot showing end-to-end PFE workflow.",
    )

    # 4. Architecture
    add_heading(doc, "4. Solution Architecture", 1)
    add_heading(doc, "4.1 High-Level Components", 2)
    add_bullets(
        doc,
        [
            "MX.3 Trade Repository — source of deals, netting, collateral agreements",
            "Market Data — curves, vol surfaces, correlation matrices for simulation",
            "Risk Engine / Simulation — PFE profile generation",
            "MLC — limit definition, risk formulae, real-time and EOD monitoring",
            "Datamart / LRB — exposure and limit utilisation reporting",
            "Workflows — pre-deal check, breach escalation, margin call validation",
        ],
    )
    add_screenshot_placeholder(
        doc,
        "MX.3 PFE / MLC architecture diagram",
        "Insert architecture diagram from your Murex PFE design doc.",
    )

    # 5. PFE Methodology
    add_heading(doc, "5. PFE Methodology", 1)
    add_heading(doc, "5.1 Definition", 2)
    add_body(
        doc,
        "PFE at confidence level α is the exposure level such that the probability of "
        "exposure exceeding PFE over the horizon is (1 − α). Exposure is typically "
        "max(MtM, 0) at each future date, aggregated at netting-set level with collateral "
        "applied per CSA terms."
    )
    add_heading(doc, "5.2 Method Options in Murex", 2)
    add_bullets(
        doc,
        [
            "Current Exposure (CE) — mark-to-market only; not PFE but often a limit bar",
            "Add-on / schedule-based — regulatory-style simplified PFE",
            "Simulation / profile-based — forward exposure paths with percentiles",
            "Monte Carlo — full portfolio simulation with netting and collateral",
        ],
    )
    add_heading(doc, "5.3 Key Parameters", 2)
    add_bullets(
        doc,
        [
            "Confidence level (e.g. 97.5%)",
            "Risk horizon (e.g. 1Y, life-to-maturity, regulatory horizon)",
            "Time buckets / profile dates",
            "Netting set and collateral eligibility rules",
            "Close-out and margin period of risk (MPOR)",
        ],
    )
    add_screenshot_placeholder(
        doc,
        "PFE simulation / profile configuration screens",
        "Insert screenshots of simulation setup, confidence level, and horizon config.",
    )

    # 6. MLC Integration
    add_heading(doc, "6. MLC Integration for PFE", 1)
    add_body(
        doc,
        "MLC (Murex Limits Controller) is the operational layer that consumes PFE (and "
        "other credit measures) and enforces limits. PFE design must align MLC static data, "
        "risk formulae, groups, and engines with the chosen PFE methodology."
    )
    add_heading(doc, "6.1 PFE in MLC Risk Types", 2)
    add_bullets(
        doc,
        [
            "Bar — single scalar PFE (e.g. max profile point or horizon PFE)",
            "Bucket — PFE by tenor / maturity bucket",
            "Profile — full exposure profile over time for monitoring and limits",
        ],
    )
    add_heading(doc, "6.2 Typical MLC Objects for PFE", 2)
    add_bullets(
        doc,
        [
            "Risk formulae referencing PFE simulation or profile outputs",
            "Limit structures mapped to counterparty / entity / product dimensions",
            "Groups and time configuration for aggregation",
            "Engines (EOD batch vs real-time) driving recalculation",
            "LRB tasks feeding Datamart exposure tables",
        ],
    )
    add_screenshot_placeholder(
        doc,
        "MLC limit structure and PFE formula configuration",
        "Insert MLC screens: limits parser, formula definition, limit assignment.",
    )

    # 7. Data & Static Setup
    add_heading(doc, "7. Data and Static Configuration", 1)
    add_bullets(
        doc,
        [
            "Counterparty static — ratings, internal grade, LEI, parent linkage",
            "Netting agreements and CSA — thresholds, MTA, eligible collateral",
            "Credit curves / funding curves assigned to counterparties",
            "Correlation matrices for multi-factor simulation (if applicable)",
            "Product / trade type mapping to risk buckets",
        ],
    )
    add_screenshot_placeholder(
        doc,
        "Static data — counterparty, CSA, credit curve assignment",
    )

    # 8. Processing
    add_heading(doc, "8. Processing — EOD and Intraday", 1)
    add_heading(doc, "8.1 EOD Sequence", 2)
    add_bullets(
        doc,
        [
            "Market data close / curve build",
            "Trade revaluation",
            "PFE simulation / profile generation",
            "MLC risk computation and limits recalculation",
            "LRB / Datamart extraction",
            "Breach reporting and workflow triggers",
        ],
    )
    add_heading(doc, "8.2 Real-Time / Intraday", 2)
    add_body(
        doc,
        "Pre-deal and amendment checks invoke MLC engines with incremental or full "
        "recalculation depending on configuration. Real-time workflows may block or warn "
        "when projected PFE exceeds limits."
    )
    add_screenshot_placeholder(
        doc,
        "EOD LTS / LRB task chain for PFE and MLC",
    )

    # 9. Reporting
    add_heading(doc, "9. Reporting and Monitoring", 1)
    add_bullets(
        doc,
        [
            "Limit utilisation by counterparty and desk",
            "PFE profile charts vs limits",
            "Top exposures and breach history",
            "Collateral coverage vs PFE",
            "Regulatory / management dashboards via Datamart",
        ],
    )
    add_screenshot_placeholder(
        doc,
        "Datamart / LRB PFE exposure report sample",
    )

    # 10. Security & Operations
    add_heading(doc, "10. Security, Operations, and Controls", 1)
    add_bullets(
        doc,
        [
            "MLC user groups and rights (view vs manage limits vs override)",
            "Audit trail for limit changes and breach comments",
            "Reconciliation between PFE engine output and MLC stored risks",
            "Failover and batch rerun procedures",
        ],
    )

    # 11. Testing
    add_heading(doc, "11. Testing Approach", 1)
    add_bullets(
        doc,
        [
            "Unit: formula and netting set logic",
            "SIT: end-to-end EOD with sample portfolios",
            "Regression: compare PFE profiles vs benchmark / prior release",
            "UAT: limit breach scenarios and pre-deal checks",
        ],
    )

    # 12. Open Points
    add_heading(doc, "12. Open Points and Assumptions", 1)
    add_bullets(
        doc,
        [
            "[ ] Confirm confidence level and horizon per limit type",
            "[ ] Confirm Monte Carlo vs simplified profile for each asset class",
            "[ ] Collateral modelling level (dynamic vs static)",
            "[ ] Real-time engine scope (products / desks)",
            "[ ] Add screenshots from source design doc into marked placeholders",
        ],
    )

    doc.add_page_break()


def build_mlc_appendix(doc):
    add_heading(doc, "Appendix A — MLC (Murex Limits Controller) Reference", 1)

    add_heading(doc, "A.1 What is MLC?", 2)
    add_body(
        doc,
        "MLC stands for Murex Limits Controller. It is the MX.3 module responsible for "
        "defining, computing, monitoring, and enforcing risk limits across credit, market, "
        "and liquidity dimensions. For counterparty credit risk, MLC is the primary "
        "operational tool used to measure exposure (including PFE, current exposure, "
        "notional, settlement risk, etc.) against approved limits."
    )
    add_body(
        doc,
        "MLC is not the PFE pricing/simulation engine itself — it orchestrates risk "
        "formulae, aggregates results, applies limit rules, and exposes monitoring and "
        "reporting. PFE values typically originate from simulation/profile engines or "
        "simplified formulae and are consumed by MLC as risk inputs."
    )

    add_heading(doc, "A.2 Where MLC Appears in a PFE Design", 2)
    add_body(doc, "In a typical Murex PFE solution design, MLC is referenced in:")
    add_bullets(
        doc,
        [
            "Architecture — MLC sits between risk computation and limit enforcement",
            "Risk types — Bar / Bucket / Profile definitions for PFE measures",
            "Limits Parser — formulae that reference PFE outputs",
            "Limit management — default limits, hierarchies, breach rules",
            "Engines — EOD risk computation and real-time pre-deal checks",
            "Workflows — Mx MLC workflows for breach handling and approvals",
            "Reporting — LRB tasks and Datamart tables for exposure vs limit",
            "Administration — user groups, rights, server properties",
        ],
    )
    add_body(
        doc,
        "Note: If your source design doc mentions MLC in sections 4 (Architecture), "
        "6 (MLC Integration), 8 (EOD processing), or Appendix configuration tables, "
        "those are the usual locations. Paste or attach those pages as screenshots "
        "into the placeholder figures in the main body."
    )

    add_heading(doc, "A.3 How MLC is Used in the Murex Credit Module", 2)
    add_bullets(
        doc,
        [
            "Pre-deal check — block or warn trades that would breach PFE limits",
            "Intraday monitoring — track limit utilisation as market moves and trades amend",
            "EOD batch — full portfolio PFE recalculation and limit re-evaluation",
            "Aggregation — roll up exposure by counterparty, entity, desk, product, tenor",
            "Netting — apply close-out netting sets and collateral per legal agreement",
            "Breach management — workflows, comments, escalations, temporary overrides",
            "Regulatory / management reporting — Datamart extracts for CRQ and dashboards",
            "Collateral linkage — exposure net of collateral where CSA is modelled",
        ],
    )

    add_heading(doc, "A.4 MLC Configuration — Step-by-Step Overview", 2)

    add_heading(doc, "A.4.1 Static Data", 3)
    add_bullets(
        doc,
        [
            "Define counterparties, entities, desks, and hierarchies used in limit dimensions",
            "Configure netting agreements and link trades to netting sets",
            "Set up CSA / collateral agreements where exposure is reduced",
            "Assign credit curves, ratings, and internal grades",
            "Define product mappings and trade attributes used in formulae (UDFs)",
        ],
    )

    add_heading(doc, "A.4.2 Limits Structure", 3)
    add_bullets(
        doc,
        [
            "Design limit hierarchy (e.g. entity → counterparty → product → bucket)",
            "Define limit types (hard block, soft warning, reporting only)",
            "Set limit amounts, currencies, and effective dates",
            "Configure default limits and inheritance rules",
        ],
    )

    add_heading(doc, "A.4.3 Groups and Time Configuration", 3)
    add_bullets(
        doc,
        [
            "Groups — define aggregation keys (counterparty group, country, sector)",
            "Time buckets — tenor bands for bucketed PFE (e.g. <1M, 1–6M, 6M–1Y, >1Y)",
            "Horizon alignment — match PFE profile dates to limit bucket definitions",
        ],
    )

    add_heading(doc, "A.4.4 Limits Parser — Risk Formulae", 3)
    add_body(
        doc,
        "The Limits Parser defines how raw trade and simulation data becomes measurable "
        "risk. For PFE, formulae typically:"
    )
    add_bullets(
        doc,
        [
            "Reference profile or bar outputs from PFE simulation",
            "Apply netting and horizontal aggregation columns",
            "Use matrices for cross-product or correlation-based rules (advanced)",
            "Define Bar (scalar), Bucket (by tenor), or Profile (time series) outputs",
        ],
    )

    add_heading(doc, "A.4.5 Engines", 3)
    add_bullets(
        doc,
        [
            "EOD engine — triggered in LTS after market close and PFE run",
            "Real-time engine — invoked on trade capture, amendment, market refresh",
            "Server properties — control recalculation scope, caching, performance",
        ],
    )

    add_heading(doc, "A.4.6 Workflows", 3)
    add_bullets(
        doc,
        [
            "Mx MLC workflows — breach notification, approval, auto-comment",
            "Pre-deal integration — hook MLC check into trade entry workflow",
            "Margin call validation — compare exposure to collateral (OSP integration)",
        ],
    )

    add_heading(doc, "A.4.7 Reporting (LRB / Datamart)", 3)
    add_bullets(
        doc,
        [
            "Create LRB (Limit Reporting Batch) tasks in LTS",
            "Configure feeders, extractions, dynamic tables, processing scripts",
            "Build Datamart reports: limit utilisation, PFE profile, breach log",
            "Schedule EOD and intraday report batches",
        ],
    )

    add_heading(doc, "A.4.8 Administration", 3)
    add_bullets(
        doc,
        [
            "User groups and rights — separate FO view, risk admin, override roles",
            "Rights templates for Datamart reports",
            "Audit and comment configuration for breaches",
        ],
    )

    add_heading(doc, "A.5 MLC vs PFE Engine — Responsibility Split", 2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Area"
    hdr[1].text = "PFE Engine / Simulation"
    hdr[2].text = "MLC"
    rows = [
        ("Compute forward exposure paths", "Yes", "No (consumes output)"),
        ("Apply confidence level / percentile", "Yes", "May reference in formula"),
        ("Define limit amounts", "No", "Yes"),
        ("Pre-deal block / warn", "Provides measure", "Enforces limit"),
        ("Datamart reporting", "Source data", "LRB extraction & reports"),
        ("Netting set logic", "Often in sim", "Aggregation in formulae"),
    ]
    for area, pfe, mlc in rows:
        row = table.add_row().cells
        row[0].text = area
        row[1].text = pfe
        row[2].text = mlc

    doc.add_paragraph()

    add_heading(doc, "A.6 Common Configuration Pitfalls", 2)
    add_bullets(
        doc,
        [
            "PFE profile dates not aligned with MLC time buckets",
            "Netting set on trades inconsistent with simulation netting",
            "Collateral not reflected in both PFE engine and MLC formula",
            "Real-time engine scope too wide — performance issues",
            "Limit currency mismatch vs exposure currency",
            "Missing LRB feeder after formula change — stale Datamart data",
        ],
    )

    add_screenshot_placeholder(
        doc,
        "MLC configuration examples from your environment",
        "Add screenshots of Limits Parser, limit assignment, and LRB task screens.",
    )


def main():
    doc = Document()
    set_document_styles(doc)
    build_main_document(doc)
    build_mlc_appendix(doc)

    output = "/workspace/Murex_PFE_Design_Document.docx"
    doc.save(output)
    print(f"Created: {output}")


if __name__ == "__main__":
    main()
